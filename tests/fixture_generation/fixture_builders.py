"""
fixture_builders.py
===================
Low-level writers for the document fixtures: EPUB (hand-built ZIP, so the
manifest order, hrefs and TOC can be made deliberately awkward), MOBI
(hand-built PalmDB/MOBI 6 records), PDF (PyMuPDF with running headers, page
numbers, a drop cap and an outline), DOCX, ODT and TXT.

Every writer is deterministic: fixed timestamps, fixed identifiers, no
randomness. Re-running the generator reproduces the same documents.
"""

from __future__ import annotations

import html
import struct
import zipfile
import zlib
from typing import Any
from urllib.parse import quote

from tests.fixture_generation import book_content as book

_FIXED_ZIP_TIME: tuple[int, int, int, int, int, int] = (2026, 1, 1, 0, 0, 0)
_FIXED_PALM_TIME: int = 0x69560000          # a fixed instant in January 2026
_XHTML_NS: str = "http://www.w3.org/1999/xhtml"
_OPS_NS: str = "http://www.idpf.org/2007/ops"


# ══════════════════════════════════════════════════════════════════════════════
# Shared helpers
# ══════════════════════════════════════════════════════════════════════════════

def make_cover_png(width: int = 120, height: int = 180) -> bytes:
    """A small deterministic PNG (striped book cover) built with zlib only."""
    rows = bytearray()
    for y in range(height):
        rows.append(0)                                  # filter type: none
        for x in range(width):
            if 20 <= y < 26 or 150 <= y < 156:
                rows.extend((240, 230, 200))            # pale bands
            elif 50 <= x < 70 and 60 <= y < 130:
                rows.extend((250, 210, 90))             # the lamp
            else:
                rows.extend((24, 48, 78 + (y * 60) // height))
    def _chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">L", len(data)) + kind + data + struct.pack(">L", zlib.crc32(kind + data) & 0xFFFFFFFF)
    header = struct.pack(">LLBBBBB", width, height, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + _chunk(b"IHDR", header)
            + _chunk(b"IDAT", zlib.compress(bytes(rows), 9)) + _chunk(b"IEND", b""))


def write_zip(path: str, members: list[tuple[str, bytes]], *, stored_first: bool = True) -> None:
    """Writes a ZIP with fixed timestamps; the first member is stored uncompressed."""
    with zipfile.ZipFile(path, "w") as archive:
        for index, (name, data) in enumerate(members):
            info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_STORED if (index == 0 and stored_first) else zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            archive.writestr(info, data)


def rezip_deterministic(path: str) -> None:
    """Rewrites a ZIP in place with fixed timestamps (python-docx/odfpy use 'now')."""
    with zipfile.ZipFile(path, "r") as archive:
        members = [(info.filename, archive.read(info.filename)) for info in archive.infolist()]
    stored_first = bool(members) and members[0][0] == "mimetype"
    write_zip(path, members, stored_first=stored_first)


def _esc(text: str) -> str:
    return html.escape(text, quote=False)


# ══════════════════════════════════════════════════════════════════════════════
# XHTML rendering
# ══════════════════════════════════════════════════════════════════════════════

def render_inline(segments: list[Any], *, note_href: str = "#fn{}") -> str:
    out = []
    for segment in segments:
        if isinstance(segment, str):
            out.append(_esc(segment))
        elif segment[0] == "i":
            out.append(f"<em>{_esc(segment[1])}</em>")
        elif segment[0] == "b":
            out.append(f"<strong>{_esc(segment[1])}</strong>")
        elif segment[0] == "br":
            out.append("<br/>")
        elif segment[0] == "fn":
            n = segment[1]
            out.append(f'<sup><a epub:type="noteref" href="{note_href.format(n)}" id="fnref{n}">{n}</a></sup>')
    return "".join(out)


def render_blocks(blocks: list[tuple], *, footnotes: bool = True) -> str:
    """Chapter body as XHTML (without the heading)."""
    out = []
    notes: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == "p":
            out.append(f"<p>{render_inline(block[1])}</p>")
            notes.extend(seg[1] for seg in block[1] if not isinstance(seg, str) and seg[0] == "fn")
        elif kind == "dropcap":
            out.append(f'<p class="first"><span class="dropcap">{block[1]}</span>{render_inline(block[2])}</p>')
        elif kind == "scene":
            out.append('<p class="scene">* * *</p>')
        elif kind == "table":
            rows = []
            for r, row in enumerate(block[1]):
                tag = "th" if r == 0 else "td"
                rows.append("<tr>" + "".join(f"<{tag}>{_esc(cell)}</{tag}>" for cell in row) + "</tr>")
            out.append("<table>" + "".join(rows) + "</table>")
    if footnotes:
        for n in notes:
            out.append(f'<aside epub:type="footnote" id="fn{n}"><p>{n}. {_esc(book.FOOTNOTES[n])}</p></aside>')
    return "\n".join(out)


def chapter_xhtml(chapter: dict[str, Any], *, heading: str = "h1", anchor: str = "") -> str:
    ident = f' id="{anchor}"' if anchor else ""
    return f"<{heading}{ident}>{_esc(chapter['title'])}</{heading}>\n{render_blocks(chapter['blocks'])}"


def title_page_xhtml() -> str:
    meta = book.TEST_BOOK_METADATA
    return f'<p class="title">{_esc(meta["title"])}</p>\n<p class="author">{_esc(meta["author"])}</p>'


def copyright_xhtml(heading: bool = False) -> str:
    head = "<h1>Copyright</h1>\n" if heading else ""
    return head + "\n".join(f"<p>{_esc(line)}</p>" for line in book.COPYRIGHT_LINES)


def about_author_xhtml(anchor: str = "") -> str:
    ident = f' id="{anchor}"' if anchor else ""
    return f"<h1{ident}>{_esc(book.ABOUT_AUTHOR_TITLE)}</h1>\n<p>{_esc(book.ABOUT_AUTHOR)}</p>"


def also_by_xhtml(anchor: str = "") -> str:
    ident = f' id="{anchor}"' if anchor else ""
    items = "".join(f"<li>{_esc(item)}</li>" for item in book.ALSO_BY)
    return f"<h1{ident}>{_esc(book.ALSO_BY_TITLE)}</h1>\n<ul>{items}</ul>"


def wrap_xhtml(title: str, body: str, body_attrs: str = "") -> bytes:
    return (
        '<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
        f'<html xmlns="{_XHTML_NS}" xmlns:epub="{_OPS_NS}" lang="en" xml:lang="en">\n'
        f"<head><title>{_esc(title)}</title></head>\n<body{body_attrs}>\n{body}\n</body>\n</html>\n"
    ).encode("utf-8")


# ══════════════════════════════════════════════════════════════════════════════
# EPUB (hand-built)
# ══════════════════════════════════════════════════════════════════════════════

def build_epub(
    path: str,
    *,
    files: list[tuple[str, str, str]],
    spine: list[str],
    toc: list[tuple[str, str]] | None,
    manifest_order: list[str] | None = None,
    nav: bool = True,
    ncx: bool = True,
    cover_png: bytes | None = None,
    encode_hrefs: bool = False,
    nonlinear: tuple[str, ...] = (),
    identifier: str = "",
    title: str = "",
) -> None:
    """Writes an EPUB from explicit parts.

    Parameters
    ----------
    path : str
        Output file.
    files : list[tuple[str, str, str]]
        (file name relative to OEBPS/, document title, XHTML body) for each content document.
    spine : list[str]
        File names in READING order.
    toc : list[tuple[str, str]] | None
        (title, href) entries; ``None`` or ``[]`` writes an empty TOC.
    manifest_order : list[str] | None
        Order of the manifest items when it should differ from the spine.
    nav, ncx : bool
        Which navigation documents to include.
    cover_png : bytes | None
        Cover image.
    encode_hrefs : bool
        Percent-encode hrefs in the manifest and the TOC (file names with
        spaces or accents), as real-world EPUBs do.
    nonlinear : tuple[str, ...]
        Spine items to mark ``linear="no"``.
    identifier, title : str
        Package metadata overrides.
    """
    meta = book.TEST_BOOK_METADATA
    identifier = identifier or meta["identifier"]
    title = title or meta["title"]
    toc = toc or []

    def _href(name: str) -> str:
        target, _, fragment = name.partition("#")
        target = quote(target) if encode_hrefs else target
        return target + (f"#{fragment}" if fragment else "")

    ids = {name: f"doc{index:02d}" for index, (name, _t, _b) in enumerate(files)}
    order = manifest_order or [name for name, _t, _b in files]
    manifest = [
        f'<item id="{ids[name]}" href="{_href(name)}" media-type="application/xhtml+xml"/>' for name in order
    ]
    if nav:
        manifest.append('<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>')
    if ncx:
        manifest.append('<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>')
    if cover_png:
        manifest.append('<item id="cover-image" href="images/cover.png" media-type="image/png" properties="cover-image"/>')
    spine_xml = "".join(
        '<itemref idref="{}"{}/>'.format(ids[name], ' linear="no"' if name in nonlinear else "")
        for name in spine
    )
    spine_attr = ' toc="ncx"' if ncx else ""
    opf = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">\n'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        f'<dc:identifier id="bookid">{_esc(identifier)}</dc:identifier>\n'
        f'<dc:title>{_esc(title)}</dc:title>\n'
        f'<dc:creator>{_esc(meta["author"])}</dc:creator>\n'
        f'<dc:language>{meta["language"]}</dc:language>\n'
        '<meta property="dcterms:modified">2026-01-01T00:00:00Z</meta>\n'
        + ('<meta name="cover" content="cover-image"/>\n' if cover_png else "")
        + "</metadata>\n<manifest>\n" + "\n".join(manifest) + "\n</manifest>\n"
        + f"<spine{spine_attr}>{spine_xml}</spine>\n</package>\n"
    )
    container = (
        '<?xml version="1.0"?>\n'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>\n'
        "</container>\n"
    )
    members: list[tuple[str, bytes]] = [
        ("mimetype", b"application/epub+zip"),
        ("META-INF/container.xml", container.encode("utf-8")),
        ("OEBPS/content.opf", opf.encode("utf-8")),
    ]
    if nav:
        items = "".join(f'<li><a href="{_esc(_href(href))}">{_esc(label)}</a></li>' for label, href in toc)
        body = f'<nav epub:type="toc" id="toc"><h1>Contents</h1><ol>{items or "<li><span>Empty</span></li>"}</ol></nav>'
        members.append(("OEBPS/nav.xhtml", wrap_xhtml("Contents", body)))
    if ncx:
        points = "".join(
            f'<navPoint id="np{index}" playOrder="{index}"><navLabel><text>{_esc(label)}</text></navLabel>'
            f'<content src="{_esc(_href(href))}"/></navPoint>'
            for index, (label, href) in enumerate(toc, start=1)
        )
        ncx_xml = (
            '<?xml version="1.0" encoding="utf-8"?>\n'
            '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">\n'
            f'<head><meta name="dtb:uid" content="{_esc(identifier)}"/></head>\n'
            f"<docTitle><text>{_esc(title)}</text></docTitle>\n<navMap>{points}</navMap>\n</ncx>\n"
        )
        members.append(("OEBPS/toc.ncx", ncx_xml.encode("utf-8")))
    for name, doc_title, body in files:
        members.append((f"OEBPS/{name}", wrap_xhtml(doc_title, body)))
    if cover_png:
        members.append(("OEBPS/images/cover.png", cover_png))
    write_zip(path, members)


# ══════════════════════════════════════════════════════════════════════════════
# TXT
# ══════════════════════════════════════════════════════════════════════════════

def build_txt(path: str) -> None:
    """Plain-text book: contents listing, underscore italics, a pipe table, ``[1]`` notes."""
    meta = book.TEST_BOOK_METADATA
    lines: list[str] = [meta["title"], f"by {meta['author']}", ""]
    lines += book.COPYRIGHT_LINES + ["", book.DEDICATION, "", "", "Contents", ""]
    lines += book.EXPECTED_TITLES + [book.NOTES_TITLE, book.ABOUT_AUTHOR_TITLE, book.ALSO_BY_TITLE, "", ""]
    for chapter in book.CHAPTERS_DATA:
        lines += [chapter["title"], ""]
        for block in chapter["blocks"]:
            if block[0] in ("p", "dropcap"):
                segments = block[1] if block[0] == "p" else [block[1] + block[2][0]] + list(block[2][1:])
                text = []
                for segment in segments:
                    if isinstance(segment, str):
                        text.append(segment)
                    elif segment[0] == "i":
                        text.append(f"_{segment[1]}_")
                    elif segment[0] == "b":
                        text.append(f"**{segment[1]}**")
                    elif segment[0] == "br":
                        text.append("\n")
                    elif segment[0] == "fn":
                        text.append(f"[{segment[1]}]")
                lines += ["".join(text), ""]
            elif block[0] == "scene":
                lines += ["* * *", ""]
            elif block[0] == "table":
                rows = block[1]
                lines.append("| " + " | ".join(rows[0]) + " |")
                lines.append("|" + "|".join("---" for _ in rows[0]) + "|")
                lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
                lines.append("")
        lines.append("")
    lines += [book.NOTES_TITLE, ""]
    lines += [f"[{n}] {text}" for n, text in book.FOOTNOTES.items()] + ["", ""]
    lines += [book.ABOUT_AUTHOR_TITLE, "", book.ABOUT_AUTHOR, "", ""]
    lines += [book.ALSO_BY_TITLE, ""] + book.ALSO_BY + [""]
    with open(path, "w", encoding="utf-8", newline="\n") as handle:
        handle.write("\n".join(lines))


# ══════════════════════════════════════════════════════════════════════════════
# DOCX
# ══════════════════════════════════════════════════════════════════════════════

_W_NS: str = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def build_docx(path: str) -> None:
    """DOCX with heading styles, a real footnote, a Word drop cap, a table and a line break."""
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import parse_xml

    meta = book.TEST_BOOK_METADATA
    document = Document()
    document.core_properties.title = meta["title"]
    document.core_properties.author = meta["author"]

    document.add_heading(meta["title"], level=0)
    document.add_paragraph(f"by {meta['author']}")
    for line in book.COPYRIGHT_LINES:
        document.add_paragraph(line)
    document.add_paragraph(book.DEDICATION)
    document.add_paragraph("Contents")
    for entry in book.EXPECTED_TITLES + [book.ABOUT_AUTHOR_TITLE, book.ALSO_BY_TITLE]:
        document.add_paragraph(entry)

    def _add_segments(paragraph, segments) -> None:
        for segment in segments:
            if isinstance(segment, str):
                paragraph.add_run(segment)
            elif segment[0] == "i":
                paragraph.add_run(segment[1]).italic = True
            elif segment[0] == "b":
                paragraph.add_run(segment[1]).bold = True
            elif segment[0] == "br":
                paragraph.add_run().add_break()
            elif segment[0] == "fn":
                paragraph._p.append(parse_xml(
                    f'<w:r xmlns:w="{_W_NS}"><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
                    f'<w:footnoteReference w:id="{int(segment[1]) + 1}"/></w:r>'
                ))

    for chapter in book.CHAPTERS_DATA:
        document.add_heading(chapter["title"], level=1)
        for block in chapter["blocks"]:
            if block[0] == "p":
                _add_segments(document.add_paragraph(), block[1])
            elif block[0] == "dropcap":
                # Word stores a drop cap as its own framed paragraph in front of the text.
                cap = document.add_paragraph()
                cap.add_run(block[1])
                cap._p.get_or_add_pPr().append(parse_xml(
                    f'<w:framePr xmlns:w="{_W_NS}" w:dropCap="drop" w:lines="2" w:wrap="around" '
                    'w:vAnchor="text" w:hAnchor="text"/>'
                ))
                _add_segments(document.add_paragraph(), block[2])
            elif block[0] == "scene":
                document.add_paragraph("* * *")
            elif block[0] == "table":
                rows = block[1]
                table = document.add_table(rows=len(rows), cols=len(rows[0]))
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        table.cell(r, c).text = cell

    document.add_heading(book.ABOUT_AUTHOR_TITLE, level=1)
    document.add_paragraph(book.ABOUT_AUTHOR)
    document.add_heading(book.ALSO_BY_TITLE, level=1)
    for item in book.ALSO_BY:
        document.add_paragraph(item)

    # A genuine footnotes part, so the marker is a real w:footnoteReference.
    notes = "".join(
        f'<w:footnote w:id="{int(n) + 1}"><w:p><w:r><w:t xml:space="preserve">{_esc(text)}</w:t></w:r></w:p></w:footnote>'
        for n, text in book.FOOTNOTES.items()
    )
    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<w:footnotes xmlns:w="{_W_NS}">'
        '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="1"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        f"{notes}</w:footnotes>"
    )
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        footnotes_xml.encode("utf-8"),
        document.part.package,
    )
    document.part.relate_to(part, RT.FOOTNOTES)

    document.save(path)
    rezip_deterministic(path)


# ══════════════════════════════════════════════════════════════════════════════
# ODT
# ══════════════════════════════════════════════════════════════════════════════

def build_odt(path: str) -> None:
    """ODT with ``text:h`` headings, a footnote, a line break, a tab and a table."""
    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import H, LineBreak, Note, NoteBody, NoteCitation, P, Span, Tab
    from odf import dc

    meta = book.TEST_BOOK_METADATA
    document = OpenDocumentText()
    document.meta.addElement(dc.Title(text=meta["title"]))
    document.meta.addElement(dc.Creator(text=meta["author"]))

    title_style = Style(name="Title", family="paragraph")
    title_style.addElement(TextProperties(fontsize="24pt", fontweight="bold"))
    document.styles.addElement(title_style)
    italic = Style(name="Emphasis", family="text")
    italic.addElement(TextProperties(fontstyle="italic"))
    document.automaticstyles.addElement(italic)
    bold = Style(name="Strong", family="text")
    bold.addElement(TextProperties(fontweight="bold"))
    document.automaticstyles.addElement(bold)

    document.text.addElement(P(stylename=title_style, text=meta["title"]))
    document.text.addElement(P(text=f"by {meta['author']}"))
    for line in book.COPYRIGHT_LINES:
        document.text.addElement(P(text=line))
    document.text.addElement(P(text=book.DEDICATION))
    document.text.addElement(P(text="Contents"))
    for entry in book.EXPECTED_TITLES + [book.ABOUT_AUTHOR_TITLE, book.ALSO_BY_TITLE]:
        # A hand-typed contents line: title, tab, page number.
        line = P(text=entry)
        line.addElement(Tab())
        line.addText("0")
        document.text.addElement(line)

    def _paragraph(segments) -> P:
        paragraph = P()
        for segment in segments:
            if isinstance(segment, str):
                paragraph.addText(segment)
            elif segment[0] == "i":
                paragraph.addElement(Span(stylename=italic, text=segment[1]))
            elif segment[0] == "b":
                paragraph.addElement(Span(stylename=bold, text=segment[1]))
            elif segment[0] == "br":
                paragraph.addElement(LineBreak())
            elif segment[0] == "fn":
                note = Note(noteclass="footnote", id=f"ftn{segment[1]}")
                note.addElement(NoteCitation(text=segment[1]))
                body = NoteBody()
                body.addElement(P(text=book.FOOTNOTES[segment[1]]))
                note.addElement(body)
                paragraph.addElement(note)
        return paragraph

    for chapter in book.CHAPTERS_DATA:
        document.text.addElement(H(outlinelevel=1, text=chapter["title"]))
        for block in chapter["blocks"]:
            if block[0] == "p":
                document.text.addElement(_paragraph(block[1]))
            elif block[0] == "dropcap":
                # ODF keeps a drop cap inline (it is a paragraph style), so the text is whole.
                document.text.addElement(_paragraph([block[1] + block[2][0]] + list(block[2][1:])))
            elif block[0] == "scene":
                document.text.addElement(P(text="* * *"))
            elif block[0] == "table":
                rows = block[1]
                table = Table(name="TideTable")
                table.addElement(TableColumn(numbercolumnsrepeated=len(rows[0])))
                for row in rows:
                    table_row = TableRow()
                    for cell in row:
                        table_cell = TableCell(valuetype="string")
                        table_cell.addElement(P(text=cell))
                        table_row.addElement(table_cell)
                    table.addElement(table_row)
                document.text.addElement(table)

    document.text.addElement(H(outlinelevel=1, text=book.ABOUT_AUTHOR_TITLE))
    document.text.addElement(P(text=book.ABOUT_AUTHOR))
    document.text.addElement(H(outlinelevel=1, text=book.ALSO_BY_TITLE))
    for item in book.ALSO_BY:
        document.text.addElement(P(text=item))

    document.save(path)
    rezip_deterministic(path)


# ══════════════════════════════════════════════════════════════════════════════
# PDF
# ══════════════════════════════════════════════════════════════════════════════

_PAGE_W: float = 420.0
_PAGE_H: float = 595.0
_MARGIN: float = 54.0
_BODY_SIZE: float = 10.5
_LEADING: float = 13.5
_BODY_TOP: float = 84.0
_BODY_BOTTOM: float = _PAGE_H - 250.0     # short pages, so chapters and paragraphs cross page breaks
_INDENT: float = 14.0
# Break points for words that may be hyphenated at a line end.
_HYPHENATION: dict[str, str] = {
    "actually": "actu-ally", "assembled": "assem-bled", "bicycle": "bicy-cle",
    "breakfast": "break-fast", "brightening": "bright-ening", "burning": "burn-ing",
    "causeway": "cause-way", "cellar": "cel-lar", "channel": "chan-nel", "cheerful": "cheer-ful",
    "children": "chil-dren", "consulted": "con-sulted", "cottage": "cot-tage", "counting": "count-ing",
    "emergency": "emer-gency", "engineer": "engi-neer", "evening": "eve-ning", "fingers": "fin-gers",
    "finished": "fin-ished", "fishing": "fish-ing", "footprints": "foot-prints",
    "generator": "gen-erator", "happened": "hap-pened", "happens": "hap-pens", "harbour": "har-bour",
    "heading": "head-ing", "imagine": "imag-ine", "impossible": "impos-sible",
    "invention": "inven-tion", "journey": "jour-ney", "language": "lan-guage", "letters": "let-ters",
    "lighthouse": "light-house", "logbook": "log-book", "master": "mas-ter", "midnight": "mid-night",
    "milestone": "mile-stone", "missing": "miss-ing", "morning": "morn-ing", "notebook": "note-book",
    "pencil": "pen-cil", "person": "per-son", "possibility": "possi-bility", "postman": "post-man",
    "promised": "prom-ised", "quarrelling": "quar-relling", "recognised": "recog-nised",
    "relight": "rel-ight", "rowing": "row-ing", "russian": "rus-sian", "salary": "sal-ary",
    "saltmarsh": "salt-marsh", "schoolmaster": "school-master", "second": "sec-ond",
    "serious": "seri-ous", "shining": "shin-ing", "simple": "sim-ple", "somebody": "some-body",
    "someone": "some-one", "stories": "sto-ries", "stubborn": "stub-born", "suggest": "sug-gest",
    "tuesday": "tues-day", "underlined": "under-lined", "understood": "under-stood",
    "visitors": "vis-itors", "watching": "watch-ing", "workshop": "work-shop", "written": "writ-ten",
}


def _import_pymupdf():
    try:
        import pymupdf
        return pymupdf
    except ImportError:
        import fitz          # PyMuPDF before 1.24.3
        return fitz


class _PdfBook:
    """A tiny line-breaking typesetter on top of PyMuPDF's ``insert_text``."""

    def __init__(self, outline: bool) -> None:
        pymupdf = _import_pymupdf()
        self.fitz = pymupdf
        self.doc = pymupdf.open()
        self.outline = outline
        self.toc: list[list] = []
        self.page = None
        self.page_number = 0
        self.y = 0.0
        self._widths: dict[tuple[str, str, float], float] = {}

    # ── measuring and drawing ──
    def _font_for(self, text: str, style: str) -> str:
        """Base-14 Times covers Latin-1 only; everything else uses a CJK system font.

        The CJK fonts are referenced, not embedded, which keeps the file small
        while the text still extracts as the right Unicode characters.
        """
        if any(0x3040 <= ord(ch) <= 0x9FFF for ch in text):
            return "japan"
        if any(ord(ch) > 0xFF for ch in text):
            return "china-s"       # Cyrillic, em-dash, ellipsis
        return {"": "tiro", "i": "tiit", "b": "tibo"}[style]

    def _width(self, text: str, font: str, size: float) -> float:
        if text.isascii():
            return self.fitz.get_text_length(text, fontname=font, fontsize=size)
        key = (text, font, size)
        if key not in self._widths:
            # Measure by drawing: the advance of non-ASCII glyphs is whatever the
            # extractor will see, which is what the layout has to agree with.
            scratch = self.fitz.open()
            page = scratch.new_page(width=3000, height=200)
            page.insert_text((10, 100), text, fontname=font, fontsize=size)
            spans = [
                span for block in page.get_text("dict")["blocks"]
                for line in block.get("lines", []) for span in line["spans"]
            ]
            self._widths[key] = (max(span["bbox"][2] for span in spans) - 10) if spans else 0.0
            scratch.close()
        return self._widths[key]

    def _draw(self, x: float, y: float, text: str, font: str, size: float) -> None:
        self.page.insert_text((x, y), text, fontname=font, fontsize=size)

    def new_page(self) -> None:
        self.page = self.doc.new_page(width=_PAGE_W, height=_PAGE_H)
        self.page_number += 1
        header = book.RUNNING_HEADER
        self._draw((_PAGE_W - self._width(header, "tiro", 8)) / 2, 34, header, "tiro", 8)
        number = str(self.page_number)
        self._draw((_PAGE_W - self._width(number, "tiro", 8)) / 2, _PAGE_H - 30, number, "tiro", 8)
        self.y = _BODY_TOP

    def _need(self, height: float) -> None:
        if self.page is None or self.y + height > _BODY_BOTTOM:
            self.new_page()

    # ── blocks ──
    def heading(self, text: str, *, size: float = 16.0, in_outline: bool = True) -> None:
        self.new_page()
        self.y = _BODY_TOP + 36
        self._draw(_MARGIN, self.y, text, "tibo", size)
        if self.outline and in_outline:
            self.toc.append([1, text, self.page_number])
        self.y += 30

    def centered(self, text: str, size: float, font: str = "tiro", gap: float = 18.0) -> None:
        self._need(gap)
        self._draw((_PAGE_W - self._width(text, font, size)) / 2, self.y, text, font, size)
        self.y += gap

    def _atoms(self, segments: list[Any]) -> list[dict]:
        """Words with their font; ``glue`` marks a piece that follows without a space."""
        atoms: list[dict] = []
        for segment in segments:
            if isinstance(segment, str):
                text, style = segment, ""
            elif segment[0] in ("i", "b"):
                text, style = segment[1], segment[0]
            elif segment[0] == "br":
                atoms.append({"break": True})
                continue
            elif segment[0] == "fn":
                atoms.append({"text": segment[1], "font": "tiro", "size": 6.5, "rise": 3.5,
                              "glue": True, "note": segment[1]})
                continue
            else:
                continue
            # Base-14 fonts cannot draw U+2026; PDFs usually spell it as three dots anyway.
            text = text.replace(chr(0x2026), "...")
            glue = (bool(atoms) and not text[:1].isspace() and "break" not in atoms[-1]
                    and not atoms[-1].get("space_after"))
            for index, word in enumerate(text.split()):
                # A word mixing scripts is drawn piece by piece, each in its own font.
                pieces: list[str] = []
                for ch in word:
                    kind = self._font_for(ch, style)
                    if pieces and self._font_for(pieces[-1][-1], style) == kind:
                        pieces[-1] += ch
                    else:
                        pieces.append(ch)
                for p, piece in enumerate(pieces):
                    atoms.append({"text": piece, "font": self._font_for(piece, style), "size": _BODY_SIZE,
                                  "rise": 0.0, "glue": (glue and index == 0 and p == 0) or p > 0})
            if text[-1:].isspace() and atoms:
                atoms[-1]["space_after"] = True
        return atoms

    def paragraph(self, segments: list[Any], *, drop_cap: str = "", indent: bool = True, size: float = _BODY_SIZE) -> None:
        atoms = self._atoms(segments)
        for atom in atoms:
            if "text" in atom and atom["size"] == _BODY_SIZE:
                atom["size"] = size
        space = self._width(" ", "tiro", size)
        right = _PAGE_W - _MARGIN
        line_no = 0
        cap_lines = 2 if drop_cap else 0
        self._need(_LEADING * (2 if drop_cap else 1))
        if drop_cap:
            self._draw(_MARGIN, self.y + _LEADING + 1, drop_cap, "tibo", 27)
        x_start = _MARGIN + (22 if cap_lines else (_INDENT if indent else 0))
        line: list[tuple[float, dict]] = []
        x = x_start
        notes: list[tuple[Any, str]] = []

        def _flush() -> None:
            nonlocal line, x, line_no, x_start
            for pos, atom in line:
                self._draw(pos, self.y - atom["rise"], atom["text"], atom["font"], atom["size"])
            line = []
            line_no += 1
            self.y += _LEADING
            if self.y > _BODY_BOTTOM:
                self.new_page()
            x_start = _MARGIN + (22 if line_no < cap_lines else 0)
            x = x_start

        pending = list(atoms)
        while pending:
            atom = pending.pop(0)
            if "break" in atom:
                _flush()
                continue
            width = self._width(atom["text"], atom["font"], atom["size"])
            lead = 0.0 if (atom["glue"] or not line) else space
            if x + lead + width > right and line:
                bare = atom["text"].strip(".,;:!?\"'")
                rule = _HYPHENATION.get(bare.lower())
                if rule and not atom["glue"]:
                    head_len = atom["text"].find(bare) + rule.index("-")
                    head, tail = atom["text"][:head_len] + "-", atom["text"][head_len:]
                    if x + lead + self._width(head, atom["font"], atom["size"]) <= right:
                        line.append((x + lead, {**atom, "text": head}))
                        pending.insert(0, {**atom, "text": tail, "glue": False})
                        _flush()
                        continue
                if atom["glue"]:
                    # Never start a line with a glued piece: take its word along.
                    carried = [line.pop()[1]]
                    while carried[0].get("glue") and line:
                        carried.insert(0, line.pop()[1])
                    pending = carried + [atom] + pending
                    _flush()
                    continue
                pending.insert(0, atom)
                _flush()
                continue
            line.append((x + lead, atom))
            x += lead + width
            if atom.get("note"):
                notes.append((self.page, atom["note"]))
        if line:
            _flush()
        for page, note in notes:
            # The footnote sits at the foot of the page that carries its marker.
            page.insert_text((_MARGIN, _PAGE_H - 62), f"{note} {book.FOOTNOTES[note]}",
                             fontname="tiro", fontsize=8)

    def table(self, rows: list[list[str]]) -> None:
        self._need(_LEADING * (len(rows) + 1))
        self.y += 4
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                self._draw(_MARGIN + 20 + c * 96, self.y, cell, "tibo" if r == 0 else "tiro", _BODY_SIZE)
            self.y += _LEADING
        self.y += 4

    def save(self, path: str) -> None:
        if self.toc:
            self.doc.set_toc(self.toc)
        meta = book.TEST_BOOK_METADATA
        self.doc.set_metadata({
            "title": meta["title"], "author": meta["author"], "producer": "AudiobookMaker fixture generator",
            "creator": "tests/fixture_generation", "creationDate": "D:20260101000000Z",
            "modDate": "D:20260101000000Z",
        })
        for page in self.doc:
            page.clean_contents()        # one content stream per page instead of one per word
        self.doc.save(path, garbage=4, deflate=True, no_new_id=True)
        self.doc.close()


def build_pdf(path: str, *, outline: bool = True) -> None:
    """PDF with a running header and a page number on EVERY page.

    Parameters
    ----------
    path : str
        Output file.
    outline : bool
        Write PDF bookmarks. Without them the chapters can only be found from
        the heading type size.
    """
    meta = book.TEST_BOOK_METADATA
    pdf = _PdfBook(outline)

    pdf.new_page()                                   # title page (never in the outline)
    pdf.y = 220
    pdf.centered(meta["title"], 22, "tibo", gap=34)
    pdf.centered(f"by {meta['author']}", 12)

    pdf.heading("Copyright")
    for line in book.COPYRIGHT_LINES:
        pdf.paragraph([line], indent=False, size=8.5)
    pdf.y += 10
    pdf.paragraph([book.DEDICATION], indent=False, size=8.5)

    pdf.heading("Contents")
    for number, entry in enumerate(book.EXPECTED_TITLES + [book.ABOUT_AUTHOR_TITLE, book.ALSO_BY_TITLE], start=4):
        pdf.paragraph([f"{entry} . . . . . {number}"], indent=False)

    for chapter in book.CHAPTERS_DATA:
        pdf.heading(chapter["title"])
        first = True
        for block in chapter["blocks"]:
            if block[0] == "p":
                pdf.paragraph(block[1], indent=not first)
            elif block[0] == "dropcap":
                pdf.paragraph(block[2], drop_cap=block[1])
            elif block[0] == "scene":
                pdf.y += 4
                pdf.centered("* * *", _BODY_SIZE, gap=_LEADING + 6)
            elif block[0] == "table":
                pdf.table(block[1])
            first = False

    pdf.heading(book.ABOUT_AUTHOR_TITLE)
    pdf.paragraph([book.ABOUT_AUTHOR], indent=False)
    pdf.heading(book.ALSO_BY_TITLE)
    for item in book.ALSO_BY:
        pdf.paragraph([item], indent=False)
    pdf.save(path)


def build_scanned_pdf(path: str, pages: int = 3) -> None:
    """A PDF with pictures but no text layer, as a flatbed scan produces."""
    pymupdf = _import_pymupdf()
    doc = pymupdf.open()
    for _ in range(pages):
        page = doc.new_page(width=_PAGE_W, height=_PAGE_H)
        for row in range(12):
            page.draw_rect(pymupdf.Rect(_MARGIN, 90 + row * 30, _PAGE_W - _MARGIN, 100 + row * 30),
                           color=(0.2, 0.2, 0.2), fill=(0.2, 0.2, 0.2))
    doc.set_metadata({"creationDate": "D:20260101000000Z", "modDate": "D:20260101000000Z"})
    doc.save(path, garbage=4, deflate=True, no_new_id=True)
    doc.close()


# ══════════════════════════════════════════════════════════════════════════════
# MOBI (hand-built PalmDB + MOBI 6 header + EXTH, uncompressed UTF-8 text)
# ══════════════════════════════════════════════════════════════════════════════

_MOBI_RECORD_SIZE: int = 4096
_NULL: int = 0xFFFFFFFF


def mobi_markup() -> bytes:
    """The book as MobiPocket markup with ``filepos`` links resolved to byte offsets."""
    meta = book.TEST_BOOK_METADATA
    pieces: list[Any] = []          # bytes | ("pos", name) | ("ref", name)

    def text(markup: str) -> None:
        pieces.append(markup.encode("utf-8"))

    def inline(segments: list[Any]) -> None:
        for segment in segments:
            if isinstance(segment, str):
                text(_esc(segment))
            elif segment[0] == "i":
                text(f"<i>{_esc(segment[1])}</i>")
            elif segment[0] == "b":
                text(f"<b>{_esc(segment[1])}</b>")
            elif segment[0] == "br":
                text("<br/>")
            elif segment[0] == "fn":
                text("<sup><a filepos=")
                pieces.append(("ref", f"note{segment[1]}"))
                text(f">{segment[1]}</a></sup>")

    text("<html><head><guide><reference type=\"toc\" title=\"Table of Contents\" filepos=")
    pieces.append(("ref", "toc"))
    text(" /></guide></head><body>")
    text(f"<h1>{_esc(meta['title'])}</h1><p>by {_esc(meta['author'])}</p><mbp:pagebreak/>")
    text("".join(f"<p>{_esc(line)}</p>" for line in book.COPYRIGHT_LINES))
    text(f"<p>{_esc(book.DEDICATION)}</p><mbp:pagebreak/>")

    pieces.append(("pos", "toc"))
    text("<h2>Contents</h2>")
    listing = [(c["title"], c["key"]) for c in book.CHAPTERS_DATA]
    listing += [(book.NOTES_TITLE, "notes"), (book.ABOUT_AUTHOR_TITLE, "about"), (book.ALSO_BY_TITLE, "alsoby")]
    for label, key in listing:
        text("<p><a filepos=")
        pieces.append(("ref", key))
        text(f">{_esc(label)}</a></p>")
    text("<mbp:pagebreak/>")

    for chapter in book.CHAPTERS_DATA:
        pieces.append(("pos", chapter["key"]))
        text(f"<h2>{_esc(chapter['title'])}</h2>")
        for block in chapter["blocks"]:
            if block[0] == "p":
                text("<p>")
                inline(block[1])
                text("</p>")
            elif block[0] == "dropcap":
                text(f"<p><font size=\"7\">{block[1]}</font>")
                inline(block[2])
                text("</p>")
            elif block[0] == "scene":
                text("<p align=\"center\">* * *</p>")
            elif block[0] == "table":
                text("<table>")
                for r, row in enumerate(block[1]):
                    tag = "th" if r == 0 else "td"
                    text("<tr>" + "".join(f"<{tag}>{_esc(cell)}</{tag}>" for cell in row) + "</tr>")
                text("</table>")
        text("<mbp:pagebreak/>")

    pieces.append(("pos", "notes"))
    text(f"<h2>{_esc(book.NOTES_TITLE)}</h2>")
    for n, note in book.FOOTNOTES.items():
        pieces.append(("pos", f"note{n}"))
        text(f"<p>{n}. {_esc(note)}</p>")
    text("<mbp:pagebreak/>")
    pieces.append(("pos", "about"))
    text(f"<h2>{_esc(book.ABOUT_AUTHOR_TITLE)}</h2><p>{_esc(book.ABOUT_AUTHOR)}</p><mbp:pagebreak/>")
    pieces.append(("pos", "alsoby"))
    text(f"<h2>{_esc(book.ALSO_BY_TITLE)}</h2>" + "".join(f"<p>{_esc(item)}</p>" for item in book.ALSO_BY))
    text("</body></html>")

    # A filepos is always written as 10 digits, so offsets are known after one pass.
    offsets: dict[str, int] = {}
    cursor = 0
    for piece in pieces:
        if isinstance(piece, bytes):
            cursor += len(piece)
        elif piece[0] == "ref":
            cursor += 10
        else:
            offsets[piece[1]] = cursor
    out = bytearray()
    for piece in pieces:
        if isinstance(piece, bytes):
            out += piece
        elif piece[0] == "ref":
            out += f"{offsets[piece[1]]:010d}".encode("ascii")
    return bytes(out)


def build_mobi(path: str, *, encrypted: bool = False) -> None:
    """Writes a classic MOBI (PalmDB container, MOBI header v6, EXTH, no compression).

    Parameters
    ----------
    path : str
        Output file.
    encrypted : bool
        Set the encryption field of the PalmDOC header, as a DRM-protected
        file has (the text itself is left alone; only the flag matters to
        readers, which must refuse the file).
    """
    meta = book.TEST_BOOK_METADATA
    markup = mobi_markup()
    text_records = [markup[i:i + _MOBI_RECORD_SIZE] for i in range(0, len(markup), _MOBI_RECORD_SIZE)]
    cover = make_cover_png()

    first_image = 1 + len(text_records)
    flis_index = first_image + 1
    fcis_index = flis_index + 1
    title = meta["title"].encode("utf-8")

    def _exth(kind: int, data: bytes) -> bytes:
        return struct.pack(">LL", kind, len(data) + 8) + data

    exth_items = [
        _exth(100, meta["author"].encode("utf-8")),
        _exth(503, title),
        _exth(524, meta["language"].encode("utf-8")),
        _exth(201, struct.pack(">L", 0)),            # cover = first image record
        _exth(202, struct.pack(">L", 0)),
    ]
    exth_body = b"".join(exth_items)
    exth = b"EXTH" + struct.pack(">LL", 12 + len(exth_body), len(exth_items)) + exth_body
    exth += b"\x00" * (-len(exth) % 4)

    header_length = 232
    palmdoc = struct.pack(">HHLHHHH", 1, 0, len(markup), len(text_records), _MOBI_RECORD_SIZE,
                          2 if encrypted else 0, 0)
    mobi = bytearray(b"MOBI")
    mobi += struct.pack(">LLLLL", header_length, 2, 65001, 0x4142_4D31, 6)   # length, type, UTF-8, uid, version
    mobi += struct.pack(">LL", _NULL, _NULL)                                 # orthographic / inflection index
    mobi += struct.pack(">LL", _NULL, _NULL)                                 # index names / keys
    mobi += struct.pack(">6L", *([_NULL] * 6))                               # extra indexes
    mobi += struct.pack(">L", first_image)                                   # first non-book record
    mobi += struct.pack(">LL", 16 + header_length + len(exth), len(title))   # full name offset / length
    mobi += struct.pack(">LLL", 0x0409, 0, 0)                                # locale, input/output language
    mobi += struct.pack(">LL", 6, first_image)                               # min version, first image record
    mobi += struct.pack(">LLLL", 0, 0, 0, 0)                                 # huffman records
    mobi += struct.pack(">L", 0x40)                                          # EXTH present
    mobi += b"\x00" * 32
    mobi += struct.pack(">L", _NULL)
    mobi += struct.pack(">LLLL", _NULL, 0, 0, 0)                             # DRM offset, count, size, flags
    mobi += b"\x00" * 8
    mobi += struct.pack(">HH", 1, len(text_records))                         # first / last content record
    mobi += struct.pack(">L", 1)
    mobi += struct.pack(">LLLL", fcis_index, 1, flis_index, 1)
    mobi += b"\x00" * 8
    mobi += struct.pack(">LLLL", _NULL, 0, _NULL, _NULL)
    mobi += struct.pack(">L", 0)                                             # extra record data flags
    mobi += struct.pack(">L", _NULL)                                         # no NCX index
    assert len(mobi) == header_length, len(mobi)
    record0 = palmdoc + bytes(mobi) + exth + title + b"\x00" * 2
    record0 += b"\x00" * (-len(record0) % 4)

    flis = b"FLIS" + struct.pack(">LHHLLHHLLL", 8, 65, 0, 0, _NULL, 1, 3, 3, 1, _NULL)
    fcis = (b"FCIS" + struct.pack(">LLLLL", 20, 16, 1, 0, len(markup))
            + struct.pack(">LLLHHL", 0, 32, 8, 1, 1, 0))
    records = [record0] + text_records + [cover, flis, fcis, b"\xe9\x8e\r\n"]

    name = b"The_Lighthouse_at_Saltmarsh"[:31]
    palm = name + b"\x00" * (32 - len(name))
    palm += struct.pack(">HHLLLLLL", 0, 0, _FIXED_PALM_TIME, _FIXED_PALM_TIME, 0, 0, 0, 0)
    palm += b"BOOKMOBI"
    palm += struct.pack(">LLH", 2 * len(records) - 1, 0, len(records))
    offset = len(palm) + 8 * len(records) + 2
    table = bytearray()
    for index, record in enumerate(records):
        table += struct.pack(">LL", offset, 2 * index)        # offset, attributes(0) + unique id
        offset += len(record)
    with open(path, "wb") as handle:
        handle.write(palm + bytes(table) + b"\x00\x00" + b"".join(records))
