"""
test_extraction_formats.py
==========================
Format-specific extraction behaviour: chapter detection in TXT, DOCX, ODT and
PDF, positional header/footer stripping and scanned-PDF handling, MOBI error
paths, and the PDF-noise step of the normaliser.

Documents are built in ``tmp_path`` with python-docx, odfpy and PyMuPDF. No
test needs Docling, EasyOCR or BookNLP.
"""

from __future__ import annotations

import os
import sys
import types

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from audiobook_factory import text_extractor  # noqa: E402
from audiobook_factory.extractor_engine import DocumentIngestor, TextNormalizer, looks_like_matter  # noqa: E402
from audiobook_factory.text_extractor import ExtractionError, extract, scan  # noqa: E402
from tests.fixture_generation import fixture_builders as build  # noqa: E402

_FIXTURES = os.path.join(_ROOT, "tests", "fixtures", "source_documents")
_PROSE = "The harbour was quiet and the tide was low that morning. "


def _quiet(_message: str) -> None:
    return None


def _flat(text: str) -> str:
    return " ".join(text.split())


def _write_txt(tmp_path, text: str, name: str = "book.txt") -> str:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return str(path)


def _titles(path: str, **kwargs) -> list[str]:
    return [c.title for c in extract(path, log_fn=_quiet, **kwargs)[0]]


def _para(marker: str, sentences: int = 12) -> str:
    return f"{marker} {_PROSE * sentences}".strip()


# ══════════════════════════════════════════════════════════════════════════════
# TXT
# ══════════════════════════════════════════════════════════════════════════════

class TestTxtChapterDetection:

    def _check(self, tmp_path, text: str, expected: list[str]) -> list:
        path = _write_txt(tmp_path, text)
        result = scan(path)
        chapters, cover = extract(path, log_fn=_quiet)
        assert cover is None
        assert [c.title for c in chapters] == expected
        assert [(c.num, c.title) for c in result.chapters] == [(c.num, c.title) for c in chapters]
        return chapters

    def test_keyword_headings_in_several_styles(self, tmp_path):
        text = "\n\n".join([
            "My Book", "Prologue", _para("ZERO"),
            "Chapter 1", _para("ONE"),
            "CHAPTER II. THE STORM", _para("TWO"),
            "Chapter Three: Home", _para("THREE"),
            "Part Two", _para("FOUR"),
            "Book III", _para("FIVE"),
            "Chapter 12 The Long Night", _para("SIX"),
            "Epilogue", _para("SEVEN"),
        ])
        chapters = self._check(tmp_path, text, [
            "Prologue", "Chapter 1", "CHAPTER II. THE STORM", "Chapter Three: Home", "Part Two", "Book III",
            "Chapter 12 The Long Night", "Epilogue",
        ])
        assert "ONE" in chapters[1].text and "TWO" not in chapters[1].text
        assert "My Book" not in " ".join(c.text for c in chapters)
        assert scan(_write_txt(tmp_path, text)).has_toc is True

    def test_gutenberg_style_stacked_and_split_titles(self, tmp_path):
        text = (
            "CHAPTER I.\nDown the Burrow\n\n" + _para("ONE") + "\n\n"
            "CHAPTER II.\nThe Pool\n\n" + _para("TWO") + "\n\n"
            "CHAPTER III\nA heading glued to its text and " + _para("THREE") + "\n"
        )
        chapters = self._check(tmp_path, text, ["CHAPTER I. Down the Burrow", "CHAPTER II. The Pool", "CHAPTER III"])
        assert "THREE" in chapters[2].text

    def test_sentences_that_start_with_chapter_words_are_not_headings(self, tmp_path):
        text = "\n\n".join([
            "Chapter 1", _para("ONE"),
            "Chapter 3 was the hardest part of her week.",
            "Part of the problem was the weather.",
            "Part one of the plan failed.",
            "Book 3 is where it all goes wrong, she said.",
            "Section 4(b), he murmured.",
            "Act now, or not at all!",
            "Prologue to a disaster, that was what it was.",
            "Introduction over, they sat down.",
            "Copyright law was her speciality.",
            "NO!",
            "\"WHAT?\"",
            _para("STILL-ONE"),
            "Chapter 2", _para("TWO"),
        ])
        chapters = self._check(tmp_path, text, ["Chapter 1", "Chapter 2"])
        assert "Chapter 3 was the hardest part of her week." in chapters[0].text
        assert "Copyright law was her speciality." in chapters[0].text
        assert "STILL-ONE" in chapters[0].text

    def test_contents_listing_does_not_create_empty_chapters(self, tmp_path):
        listing = ["Chapter 1 ........ 3", "Chapter 2 ........ 9", "Chapter 3 ........ 15", "About the Author .... 21"]
        text = "\n\n".join(["A Book", "CONTENTS", "\n".join(listing),
                           "Chapter 1", _para("ONE"), "Chapter 2", _para("TWO"), "Chapter 3", _para("THREE"),
                           "About the Author", "The author lives by the sea."])
        chapters = self._check(tmp_path, text, ["Chapter 1", "Chapter 2", "Chapter 3"])
        spoken = " ".join(c.text for c in chapters)
        assert "........" not in spoken and "lives by the sea" not in spoken and "CONTENTS" not in spoken
        flagged = [c.title for c in scan(_write_txt(tmp_path, text), include_matter=True).chapters if c.probably_matter]
        assert "CONTENTS" in flagged and "About the Author" in flagged

    def test_unheaded_contents_run_is_dropped(self, tmp_path):
        text = "\n\n".join(["A Book", "Chapter 1", "Chapter 2", "Chapter 3",
                            "Chapter 1", _para("ONE"), "Chapter 2", _para("TWO"), "Chapter 3", _para("THREE")])
        chapters = self._check(tmp_path, text, ["Chapter 1", "Chapter 2", "Chapter 3"])
        assert all(len(c.text) > 200 for c in chapters)
        assert chapters[0].text.count("Chapter 1") == 1

    def test_contents_listing_finds_plain_titles(self, tmp_path):
        text = "\n\n".join(["Contents", "The Salt Road\nAbout the Lighthouse\nHome Again",
                            "The Salt Road", _para("ONE"), "About the Lighthouse", _para("TWO"),
                            "Home Again", _para("THREE")])
        self._check(tmp_path, text, ["The Salt Road", "About the Lighthouse", "Home Again"])

    def test_bare_numbers_and_roman_numerals(self, tmp_path):
        numbered = "\n\n".join(part for n in range(1, 5) for part in (str(n), _para(f"N{n}", 70)))
        self._check(tmp_path, numbered, ["1", "2", "3", "4"])
        roman = "\n\n".join(part for n, r in enumerate(["I.", "II.", "III."]) for part in (r, _para(f"R{n}", 70)))
        self._check(tmp_path, roman, ["I", "II", "III"])

    def test_page_numbers_are_not_chapters(self, tmp_path):
        pages = "\n\n".join(part for n in range(1, 9) for part in (str(n), _para(f"P{n}", 8)))
        path = _write_txt(tmp_path, pages)
        assert _titles(path) == ["Full Book"]
        assert scan(path).has_toc is False

    def test_all_caps_headings(self, tmp_path):
        text = "\n\n".join(["THE ARRIVAL", _para("ONE", 40), "A SHORT LINE OF SHOUTING!", _para("MORE", 5),
                            "THE STORM", _para("TWO", 40), "HOME", _para("THREE", 40)])
        self._check(tmp_path, text, ["THE ARRIVAL", "THE STORM", "HOME"])

    def test_markdown_headings_and_structure(self, tmp_path):
        text = ("# The Book\n\n## One\n\n" + _para("ONE") + "\n\n- first item\n- second item\n\n"
                "| Day | High |\n|---|---|\n| Monday | 6:10 |\n\n## Two\n\n" + _para("TWO")
                + " See [the map](http://example.com/m).[3]\n\n[3] A footnote that is not read.\n")
        chapters = self._check(tmp_path, text, ["One", "Two"])
        assert chapters[0].text.startswith("One\n\nONE")
        assert "first item." in chapters[0].text and "Monday, 6:10." in chapters[0].text
        spoken = " ".join(c.text for c in chapters)
        for junk in ("#", "|", "](", "http", "[3]", "footnote that is not read", "- first"):
            assert junk not in spoken

    def test_text_without_headings_is_one_full_book_chapter(self, tmp_path):
        path = _write_txt(tmp_path, "Just a short note.\n\nSecond paragraph of the note.")
        result = scan(path)
        assert result.has_toc is False and result.page_count == 0 and not result.supports_page_ranges
        assert [(c.num, c.title) for c in result.chapters] == [(1, "Full Book")]
        chapters, _ = extract(path, log_fn=_quiet)
        assert [c.title for c in chapters] == ["Full Book"]
        assert chapters[0].text == "Just a short note.\n\nSecond paragraph of the note."

    def test_project_gutenberg_wrapper_is_not_narrated(self, tmp_path):
        text = "\n\n".join([
            "The Project Gutenberg eBook of A Tale", "This ebook is for the use of anyone anywhere.",
            "*** START OF THE PROJECT GUTENBERG EBOOK A TALE ***",
            _para("STORY", 30),
            "*** END OF THE PROJECT GUTENBERG EBOOK A TALE ***",
            "Section 1. General Terms of Use and Redistributing Project Gutenberg-tm electronic works " * 5,
        ])
        path = _write_txt(tmp_path, text)
        chapters, _ = extract(path, log_fn=_quiet)
        assert [c.title for c in chapters] == ["Full Book"]
        assert "STORY" in chapters[0].text and "Gutenberg" not in chapters[0].text
        flagged = scan(path, include_matter=True).chapters
        assert [c.title for c in flagged if c.probably_matter] == [
            "Project Gutenberg Header", "Project Gutenberg License"]

    def test_long_dedication_heading_does_not_swallow_the_book(self, tmp_path):
        text = "\n\n".join(["Dedication", "For Mum."] + [_para(f"STORY{n}", 30) for n in range(12)])
        chapters, _ = extract(_write_txt(tmp_path, text), log_fn=_quiet)
        assert len(chapters) == 1 and "STORY11" in chapters[0].text

    def test_selection_and_legacy_encoding(self, tmp_path):
        text = "\n\n".join(["Chapter 1", _para("It’s a café."), "Chapter 2", _para("TWO")])
        path = tmp_path / "legacy.txt"
        path.write_bytes(text.encode("cp1252"))
        chapters, _ = extract(str(path), selections=["Chapter 1"], log_fn=_quiet)
        assert [c.title for c in chapters] == ["Chapter 1"]
        assert "It's a café." in chapters[0].text


# ══════════════════════════════════════════════════════════════════════════════
# DOCX
# ══════════════════════════════════════════════════════════════════════════════

class TestDocx:

    def _document(self):
        docx = pytest.importorskip("docx")
        return docx.Document()

    def test_tables_are_read_in_document_order(self, tmp_path):
        document = self._document()
        document.add_heading("Chapter 1", level=1)
        document.add_paragraph("Before the table. " + _PROSE * 4)
        table = document.add_table(rows=2, cols=2)
        for (r, c), value in {(0, 0): "Item", (0, 1): "Price", (1, 0): "Tea", (1, 1): "3 shillings"}.items():
            table.cell(r, c).text = value
        document.add_paragraph("After the table. " + _PROSE * 4)
        path = str(tmp_path / "table.docx")
        document.save(path)
        text = extract(path, log_fn=_quiet)[0][0].text
        assert text.index("Before the table.") < text.index("Item, Price") < text.index("Tea, 3 shillings") \
            < text.index("After the table.")

    def test_heading_levels_and_parts(self, tmp_path):
        document = self._document()
        document.add_heading("The Book", level=0)
        for part, chapters in (("Part One", ("Dawn", "Noon")), ("Part Two", ("Dusk",))):
            document.add_heading(part, level=1)
            for title in chapters:
                document.add_heading(title, level=2)
                document.add_paragraph(_para(title.upper()))
                document.add_heading("A sub-section", level=3)
                document.add_paragraph(_para("SUB"))
        path = str(tmp_path / "parts.docx")
        document.save(path)
        chapters, _ = extract(path, log_fn=_quiet)
        assert [c.title for c in chapters] == ["Dawn", "Noon", "Dusk"]
        assert "Part One" in chapters[0].text and "Part Two" in chapters[2].text    # nothing is lost
        assert "A sub-section" in chapters[1].text
        assert [c.title for c in scan(path).chapters] == ["Dawn", "Noon", "Dusk"]

    def test_tracked_changes_toc_lines_and_field_codes(self, tmp_path):
        from docx.oxml import parse_xml
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        document = self._document()
        document.add_heading("Chapter 1", level=1)
        paragraph = document.add_paragraph("The tide was ")
        paragraph._p.append(parse_xml(
            f'<w:del xmlns:w="{ns}" w:id="1" w:author="a"><w:r><w:delText>low</w:delText></w:r></w:del>'))
        paragraph._p.append(parse_xml(
            f'<w:ins xmlns:w="{ns}" w:id="2" w:author="a"><w:r><w:t>high</w:t></w:r></w:ins>'))
        paragraph.add_run(" that day.")
        paragraph._p.append(parse_xml(
            f'<w:r xmlns:w="{ns}"><w:instrText> PAGEREF _Toc1 </w:instrText></w:r>'))
        tabbed = document.add_paragraph("Low water:")
        tabbed.add_run().add_tab()
        tabbed.add_run("12:25")
        document.add_paragraph("Chapter 9 ...... 99", style="TOC Heading")
        document.add_paragraph(_PROSE * 5)
        path = str(tmp_path / "tracked.docx")
        document.save(path)
        text = extract(path, log_fn=_quiet)[0][0].text
        assert "The tide was high that day." in text
        assert "Low water: 12:25" in text
        assert "PAGEREF" not in text and "......" not in text and "low" not in text.split("that day.")[0]

    def test_document_without_headings_is_one_chapter_and_has_no_pages(self, tmp_path):
        document = self._document()
        document.add_paragraph("Just one paragraph of text in a plain document.")
        path = str(tmp_path / "plain.docx")
        document.save(path)
        result = scan(path)
        assert (result.has_toc, result.page_count, result.supports_page_ranges) == (False, 0, False)
        assert [c.title for c in result.chapters] == ["Full Book"]
        chapters, _ = extract(path, page_ranges=[(1, 2)], log_fn=_quiet)     # ranges mean nothing for DOCX
        assert [c.title for c in chapters] == ["Full Book"]

    def test_fixture_drop_cap_footnote_and_line_break(self):
        chapters, _ = extract(os.path.join(_FIXTURES, "dummy_book.docx"), log_fn=_quiet)
        prologue, chapter_one = chapters[0].text, chapters[1].text
        assert "The tide went out at dawn" in prologue and "T he" not in prologue and "\nT\n" not in prologue
        assert "tide tables twice before breakfast" in chapter_one
        assert "read:\nKEEPER ABSENT\nENQUIRE AT THE HARBOUR OFFICE" in chapter_one


# ══════════════════════════════════════════════════════════════════════════════
# ODT
# ══════════════════════════════════════════════════════════════════════════════

class TestOdt:

    def _build(self, tmp_path) -> str:
        pytest.importorskip("odf")
        from odf.opendocument import OpenDocumentText
        from odf.table import Table, TableCell, TableColumn, TableRow
        from odf.text import H, LineBreak, List, ListItem, Note, NoteBody, NoteCitation, P, S, Tab, TrackedChanges

        document = OpenDocumentText()
        body = document.text
        changes = TrackedChanges()
        changes.addElement(P(text="DELETED SENTENCE THAT WAS TRACKED"), check_grammar=False)
        body.addElement(changes)
        body.addElement(H(outlinelevel=1, text="Chapter One"))
        paragraph = P(text="First line")
        paragraph.addElement(LineBreak())
        paragraph.addText("second line")
        paragraph.addElement(Tab())
        paragraph.addText("after tab")
        paragraph.addElement(S(c=2))
        paragraph.addText("after spaces, with a note")
        note = Note(noteclass="footnote", id="ftn1")
        note.addElement(NoteCitation(text="1"))
        note_body = NoteBody()
        note_body.addElement(P(text="FOOTNOTE BODY TEXT"))
        note.addElement(note_body)
        paragraph.addElement(note)
        paragraph.addText(" in the middle of it. " + _PROSE * 5)
        body.addElement(paragraph)
        items = List()
        for label in ("List item one.", "List item two."):
            item = ListItem()
            item.addElement(P(text=label))
            items.addElement(item)
        body.addElement(items)
        table = Table(name="T")
        table.addElement(TableColumn(numbercolumnsrepeated=2))
        for row in (("Item", "Price"), ("Tea", "3 shillings")):
            table_row = TableRow()
            for value in row:
                cell = TableCell(valuetype="string")
                cell.addElement(P(text=value))
                table_row.addElement(cell)
            table.addElement(table_row)
        body.addElement(table)
        body.addElement(H(outlinelevel=1, text="Chapter Two"))
        body.addElement(P(text=_para("TWO")))
        path = str(tmp_path / "book.odt")
        document.save(path)
        return path

    def test_headings_breaks_notes_lists_and_tables(self, tmp_path):
        path = self._build(tmp_path)
        chapters, _ = extract(path, log_fn=_quiet)
        assert [c.title for c in chapters] == ["Chapter One", "Chapter Two"]
        assert [c.title for c in scan(path).chapters] == ["Chapter One", "Chapter Two"]
        text = chapters[0].text
        # A line break and a tab separate words instead of gluing them together.
        assert "First line second line after tab after spaces, with a note" in _flat(text)
        assert "with a note in the middle of it." in text
        assert "FOOTNOTE BODY TEXT" not in text and "DELETED SENTENCE" not in text
        assert text.count("in the middle of it.") == 1
        assert "List item one.\n\nList item two." in text
        assert text.index("List item two.") < text.index("Item, Price") < text.index("Tea, 3 shillings")
        assert scan(path).page_count == 0


# ══════════════════════════════════════════════════════════════════════════════
# PDF
# ══════════════════════════════════════════════════════════════════════════════

def _make_pdf(tmp_path, pages: list[dict], name: str = "book.pdf", toc: list | None = None) -> str:
    """pages: dicts with optional 'header', 'footer', 'heading' and 'lines'."""
    pymupdf = pytest.importorskip("pymupdf")
    doc = pymupdf.open()
    for spec in pages:
        page = doc.new_page(width=420, height=595)
        if spec.get("header"):
            page.insert_text((54, 34), spec["header"], fontname="tiro", fontsize=9)
        if spec.get("footer"):
            page.insert_text((200, 565), spec["footer"], fontname="tiro", fontsize=9)
        y = 110.0
        if spec.get("heading"):
            page.insert_text((54, y), spec["heading"], fontname="tibo", fontsize=spec.get("heading_size", 16))
            y += 34
        for line in spec.get("lines", []):
            if line == "":
                y += 10
                continue
            page.insert_text((54, y), line, fontname="tiro", fontsize=10.5)
            y += 13.5
    if toc:
        doc.set_toc(toc)
    path = str(tmp_path / name)
    doc.save(path)
    doc.close()
    return path


def _body_lines(marker: str, count: int = 8) -> list[str]:
    return [f"{marker} line {n} of the page keeps the story going along nicely." for n in range(count)]


class TestPdf:

    def test_header_with_page_number_and_repeated_dialogue(self, tmp_path):
        pages = []
        for number in range(1, 9):
            lines = _body_lines(f"P{number}", 4) + ["", "\"What?\"", "", "Introduction"] + _body_lines(f"Q{number}", 2)
            pages.append({"header": f"{number}   A Tale of Tides", "footer": f"Page {number} of 8", "lines": lines})
        path = _make_pdf(tmp_path, pages)
        chapters, _ = extract(path, log_fn=_quiet)
        text = "\n".join(c.text for c in chapters)
        assert "A Tale of Tides" not in text and "Page " not in text
        assert text.count("\"What?\"") == 8, "repeated dialogue must survive"
        assert text.count("Introduction") == 8
        assert "P1 line 0" in text and "Q8 line 1" in text

    def test_outline_chapters_and_same_page_split(self, tmp_path):
        pages = [
            {"header": "A Tale", "footer": "1", "lines": ["A Tale", "by Nobody"]},
            {"header": "A Tale", "footer": "2", "heading": "Dawn", "lines": _body_lines("DAWN", 6)},
            {"header": "A Tale", "footer": "3", "lines": _body_lines("DAWNTWO", 4) + ["", "Noon", ""] + _body_lines("NOON", 6)},
            {"header": "A Tale", "footer": "4", "heading": "Dusk", "lines": _body_lines("DUSK", 6)},
        ]
        path = _make_pdf(tmp_path, pages, toc=[[1, "Dawn", 2], [1, "Noon", 3], [1, "Dusk", 4]])
        result = scan(path)
        assert (result.page_count, result.supports_page_ranges, result.has_toc) == (4, True, True)
        chapters, _ = extract(path, log_fn=_quiet)
        assert [c.title for c in chapters] == ["Dawn", "Noon", "Dusk"]
        assert [c.title for c in result.chapters] == ["Dawn", "Noon", "Dusk"]
        assert "DAWNTWO line 3" in chapters[0].text and "NOON" not in chapters[0].text
        assert chapters[1].text.startswith("Noon") and "NOON line 5" in chapters[1].text
        only, _ = extract(path, selections=["Noon"], log_fn=_quiet)
        assert [c.title for c in only] == ["Noon"]

    def test_keyword_headings_without_outline_or_large_type(self, tmp_path):
        pages = [{"header": "A Tale", "footer": str(n + 1),
                  "lines": [f"Chapter {n + 1}", ""] + _body_lines(f"C{n + 1}", 10)} for n in range(3)]
        path = _make_pdf(tmp_path, pages)
        assert _titles(path) == ["Chapter 1", "Chapter 2", "Chapter 3"]

    def test_pdf_without_structure_is_one_chapter_and_offers_page_ranges(self, tmp_path):
        path = _make_pdf(tmp_path, [{"lines": _body_lines("ONLY", 5)}, {"lines": _body_lines("MORE", 5)}])
        result = scan(path)
        assert (result.has_toc, result.page_count, result.supports_page_ranges) == (False, 2, True)
        assert [c.title for c in result.chapters] == ["Full Book"]
        assert _titles(path) == ["Full Book"]
        assert _titles(path, page_ranges=[(2, 2)]) == ["Chapter 1 (pp. 2–2)"]

    def test_scanned_pdf_is_reported_clearly(self, tmp_path):
        path = str(tmp_path / "scan.pdf")
        build.build_scanned_pdf(path)
        result = scan(path)
        assert result.page_count == 3 and not result.chapters and "no text layer" in result.warning
        with pytest.raises(ExtractionError, match="no text layer"):
            extract(path, log_fn=_quiet)
        with pytest.raises(ExtractionError, match="no text layer"):
            extract(path, page_ranges=[(1, 2)], log_fn=_quiet)
        assert issubclass(ExtractionError, ValueError)
        import importlib.util
        if importlib.util.find_spec("easyocr") is None:
            # OCR was asked for but no OCR engine is installed: same clear error.
            with pytest.raises(ExtractionError, match="no text layer"):
                extract(path, enable_ocr=True, log_fn=_quiet)

    def test_scanned_pdf_is_ocrd_when_asked_and_possible(self, tmp_path, monkeypatch):
        pytest.importorskip("numpy")
        readers = []

        class _Reader:
            def __init__(self, languages, gpu=True):
                self.languages, self.gpu = languages, gpu
                readers.append(self)

            def readtext(self, image, **kwargs):
                assert image.ndim == 3 and image.shape[2] == 3
                return ["SCANNED PAGE TEXT that was only a picture."]

        monkeypatch.setitem(sys.modules, "easyocr", types.SimpleNamespace(Reader=_Reader))
        monkeypatch.delenv("AUDIOBOOK_OCR_GPU", raising=False)
        DocumentIngestor.release_ocr_reader()
        path = str(tmp_path / "scan.pdf")
        build.build_scanned_pdf(path, pages=2)
        chapters, _ = extract(path, enable_ocr=True, language="de", log_fn=_quiet)
        assert [c.title for c in chapters] == ["Full Book"]
        assert chapters[0].text.count("SCANNED PAGE TEXT") == 2
        assert [(r.languages, r.gpu) for r in readers] == [(["de", "en"], False)]
        assert DocumentIngestor._easyocr_reader is None

    def test_document_is_closed_even_when_extraction_fails(self, tmp_path, monkeypatch):
        import fitz
        path = str(tmp_path / "scan.pdf")
        build.build_scanned_pdf(path)
        opened = []
        real_open = fitz.open

        def _open(*args, **kwargs):
            opened.append(real_open(*args, **kwargs))
            return opened[-1]

        monkeypatch.setattr(fitz, "open", _open)
        with pytest.raises(ExtractionError):
            extract(path, log_fn=_quiet)
        scan(path)
        assert len(opened) == 2 and all(doc.is_closed for doc in opened)

    def test_fixture_pdf_paragraphs_cross_page_breaks(self):
        import fitz
        path = os.path.join(_FIXTURES, "dummy_book.pdf")
        doc = fitz.open(path)
        try:
            mid_sentence = 0
            for page in doc:
                lines = [ln for ln in page.get_text("text").split("\n") if ln.strip()]
                body = [ln for ln in lines if ln.strip() not in ("The Lighthouse at Saltmarsh Point",)
                        and not ln.strip().isdigit()]
                if body and body[-1].rstrip()[-1:].isalpha():
                    mid_sentence += 1
        finally:
            doc.close()
        assert mid_sentence >= 2, "the fixture should break pages in the middle of a paragraph"
        chapter_one = extract(path, selections=["Chapter 1: The Salt Road"], log_fn=_quiet)[0][0]
        assert "e.g., a broken lens, a flooded cellar, a failed generator, etc., and for none" in chapter_one.text


# ══════════════════════════════════════════════════════════════════════════════
# MOBI error paths
# ══════════════════════════════════════════════════════════════════════════════

class TestMobiErrors:

    def test_drm_protected_file_gives_a_clear_error(self, tmp_path):
        path = str(tmp_path / "locked.mobi")
        build.build_mobi(path, encrypted=True)
        with pytest.raises(ExtractionError, match="DRM"):
            extract(path, log_fn=_quiet)
        result = scan(path)
        assert result.file_type == "mobi" and not result.chapters and "DRM" in result.warning

    def test_file_that_is_not_a_mobi(self, tmp_path):
        path = tmp_path / "fake.mobi"
        with open(os.path.join(_FIXTURES, "dummy_book.epub"), "rb") as handle:
            path.write_bytes(handle.read())
        with pytest.raises(ExtractionError, match="Not a valid MOBI"):
            extract(str(path), log_fn=_quiet)
        assert "Not a valid MOBI" in scan(str(path)).warning
        empty = tmp_path / "empty.azw3"
        empty.write_bytes(b"")
        with pytest.raises(ExtractionError):
            extract(str(empty), log_fn=_quiet)

    def test_missing_mobi_package_gives_an_install_hint(self, monkeypatch):
        path = os.path.join(_FIXTURES, "dummy_book.mobi")
        monkeypatch.setitem(sys.modules, "mobi", None)
        monkeypatch.setitem(sys.modules, "mobi.kindleunpack", None)
        with pytest.raises(ExtractionError, match="pip install mobi"):
            extract(path, log_fn=_quiet)
        assert "pip install mobi" in scan(path).warning

    def test_corrupt_records_do_not_crash(self, tmp_path):
        pytest.importorskip("mobi")
        with open(os.path.join(_FIXTURES, "dummy_book.mobi"), "rb") as handle:
            data = bytearray(handle.read())
        del data[len(data) // 3:]                       # truncated download
        path = tmp_path / "broken.mobi"
        path.write_bytes(bytes(data))
        try:
            extract(str(path), log_fn=_quiet)
        except ExtractionError as error:
            assert "MOBI" in str(error)
        assert scan(str(path)).file_type == "mobi"

    @pytest.mark.parametrize("name", ["book.mobi", "book.azw3", "book.azw", "book.prc", "BOOK.MOBI"])
    def test_kindle_extensions_are_recognised(self, name):
        assert text_extractor._detect_type(name) == "mobi"

    def test_classic_mobi_without_contents_page_is_split_at_headings(self, tmp_path):
        """Unpacked MOBI 7 markup with no listing: links in prose stay prose."""
        markup = (
            "<html><head></head><body><h1>A Tale</h1><p>by Nobody</p><mbp:pagebreak/>"
            "<h2>Dawn</h2><p>ALPHA she read the first note"
            '<sup><a href="#filepos900">1</a></sup> and KEEPME the second note'
            '<sup><a href="#filepos950">2</a></sup> before breakfast. ' + _PROSE * 6 + "</p><mbp:pagebreak/>"
            "<h2>Dusk</h2><p>BETA " + _PROSE * 6 + ' See <a href="#filepos10">the start</a>.</p>'
            "</body></html>"
        )
        (tmp_path / "book.html").write_text(markup, encoding="utf-8")
        _ingestor, sections, title, author, _language, cover = text_extractor._mobi7_plan(str(tmp_path / "book.html"))
        assert [(s.title, s.probably_matter) for s in sections if not s.probably_matter] == [
            ("Dawn", False), ("Dusk", False)]
        dawn = "".join(fragment for _doc, fragment in sections[[s.title for s in sections].index("Dawn")].parts)
        assert "KEEPME the second note" in dawn
        assert (title, author, cover) == ("", "", None)


# ══════════════════════════════════════════════════════════════════════════════
# Normaliser: PDF noise step
# ══════════════════════════════════════════════════════════════════════════════

class TestPdfNoiseStep:

    RAW = "\n".join([
        "## Introduction", "", "A Tale of Tides", "", "The first page of the introduction ends here.", "",
        "12", "", "A Tale of Tides", "", "\"What?\"", "", "She asked again.", "", "\"What?\"", "",
        "13", "", "A Tale of Tides", "", "Introduction", "", "\"What?\"", "", "Introduction", "",
        "<!-- image -->", "", "Introduction", "", "ThisIsGarbledOcrSoup", "", "The end of it.",
    ])

    def _normalize(self, **kwargs) -> str:
        return TextNormalizer().normalize(self.RAW, "A Tale", [], **kwargs)

    def test_repeated_dialogue_and_headings_survive(self):
        text = self._normalize(is_pdf=True)
        lines = [line for line in text.split("\n") if line]
        assert lines.count("\"What?\"") == 3
        assert lines[0] == "Introduction"                     # the real heading
        assert lines.count("Introduction") == 2               # heading + first bare occurrence
        assert lines.count("A Tale of Tides") == 1            # running header: first occurrence only
        assert "12" not in lines and "13" not in lines
        assert "ThisIsGarbledOcrSoup" not in text and "image" not in text
        assert "She asked again." in text and "The end of it." in text

    def test_without_the_pdf_flag_nothing_is_removed(self):
        text = self._normalize()
        assert text.count("A Tale of Tides") == 3 and "12" in text.split("\n")

    def test_rust_and_python_agree_on_pdf_text(self, monkeypatch):
        pytest.importorskip("audiobook_rust")
        from audiobook_factory import text_processing
        with_rust = [self._normalize(is_pdf=True), self._normalize(fix_kerning=True), self._normalize()]
        monkeypatch.setitem(sys.modules, "audiobook_rust", None)
        monkeypatch.setattr(text_processing, "_RUST_AVAILABLE", False)
        assert [self._normalize(is_pdf=True), self._normalize(fix_kerning=True), self._normalize()] == with_rust

    def test_kerning_repair_still_runs_for_docling_pdf_text(self):
        fixed = TextNormalizer().normalize("W ar came to T ohsaka. Class D students left.", "T", [], is_pdf=True)
        assert fixed == "War came to Tohsaka. Class D students left."

    @pytest.mark.parametrize("text, expected", [
        ("Copyright © 2026 Mara Ellison. All rights reserved.", True),
        ("ISBN 978-0-00-000000-0\nFirst published by Test Press.", True),
        ("Chapter 1 ..... 3\nChapter 2 ..... 9\nChapter 3 ..... 15\nEpilogue ..... 40", True),
        ("First published by Test Press. Printed in the United Kingdom.", True),
        ("She found the first edition of the paper on the table and read it twice over.", False),
        ("The tide went out at dawn, as it had every morning since 1887.", False),
        ("\"No!\" he cried.\n\"Yes,\" she said.\n\"Why?\"\n\"Because.\"\nHe left.", False),
    ])
    def test_looks_like_matter(self, text, expected):
        assert looks_like_matter(text) is expected
